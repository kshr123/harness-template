"""WBS を表計算（.xlsx）に書き出す。**先方の様式指定に応えるための枝**であって、既定の成果物ではない。

日本の受託では、先方の管理部門が「WBS は Excel で」と様式を指定してくることが実際にある。HTML しか
出せないと、そこで折れる。ただし**取り込みは作らない**：表計算は編集を誘う道具なので、返送された
ファイルを読み込む口を開けると、行の挿入・並べ替え・書式でセルの対応が黙って壊れる（＝二重台帳への入口）。
このファイルは読み取り専用の写しであることを、シートの先頭に明記して渡す。

openpyxl が入っている案件でだけ登録される（`formats.py` の条件登録）。核はこの形式を知らない。
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path
from typing import TYPE_CHECKING, Any

from harness.deliver.render import COLUMNS as _COLUMNS  # 列の並びと表示名は 1 か所（形式ごとに写しを持たない）
from harness.deliver.render import STATUS_LABEL as _STATUS_LABEL
from harness.deliver.wbs import Wbs, WbsRow
from harness.models import Status

if TYPE_CHECKING:  # 型だけ（実体は書き出すときに取り込む）
    from openpyxl.worksheet.worksheet import Worksheet

# 週の列の塗り（HTML の棒と同じ意味・同じ色合い）。
_FILL = {"plan": "5B87B8", "done": "4F9D72", "late": "C8635A", "ms": "1B1F24"}

_NOTICE = "この表は生成した写しです。正本は work/ の作業単位で、この表への記入は取り込まれません。"


def _weeks(span: tuple[date, date]) -> list[date]:
    """期間を覆う週（月曜）の並び。列 1 つが 1 週間。"""
    first, last = span
    monday = first - timedelta(days=first.weekday())
    out: list[date] = []
    while monday <= last:
        out.append(monday)
        monday += timedelta(days=7)
    return out


def _kind_of(row: WbsRow) -> str | None:
    """その行を週の列でどう塗るか（塗らないなら None）。"""
    if row.milestone and row.due is not None:
        return "ms"
    if row.start is None or row.due is None:
        return None
    if row.late:
        return "late"
    return "done" if row.status is Status.done else "plan"


def _write_row(sheet: Worksheet, index: int, row: WbsRow, weeks: list[date], fills: dict[str, Any]) -> None:
    """1 行を書く（左の表＋右の週ごとの塗り）。字下げと折りたたみは行の階層で表す。"""
    depth = row.code.count(".")
    values: list[str | float | date | None] = [
        row.code,
        ("　" * depth) + row.name,
        row.team or "",
        "、".join(row.assignees),
        _STATUS_LABEL[row.status] if row.status is not None else "",
        row.start,
        row.due,
        row.workdays,
        row.actual_start,
        row.actual_finish,
        f"{row.done_leaves}/{row.total_leaves}" if row.total_leaves else "",
    ]
    for column, value in enumerate(values, start=1):
        sheet.cell(row=index, column=column, value=value)
    # 表計算側の折りたたみ（アウトライン）。第 1 階層は 0 なので、そのまま深さを使う。
    sheet.row_dimensions[index].outlineLevel = depth
    kind = _kind_of(row)
    if kind is None:
        return
    start = row.due if row.milestone else row.start
    end = row.due
    if start is None or end is None:
        return
    for offset, monday in enumerate(weeks):
        if monday + timedelta(days=6) >= start and monday <= end:
            sheet.cell(row=index, column=len(_COLUMNS) + 1 + offset).fill = fills[kind]


def write_xlsx(wbs: Wbs, path: Path, *, provenance: str = "", draft: bool = False) -> None:
    """表計算の写し（先方の様式指定がある案件向け。取り込みは作らない）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    span = wbs.span
    weeks = _weeks(span) if span else []
    book = Workbook()
    sheet = book.active
    if sheet is None:  # pragma: no cover - openpyxl は必ず 1 枚目を作る
        sheet = book.create_sheet()
    sheet.title = "WBS"

    sheet.cell(row=1, column=1, value=wbs.overlay.project or "WBS").font = Font(bold=True, size=13)
    sheet.cell(row=2, column=1, value=f"基準日 {wbs.today.isoformat()}　{provenance}")
    sheet.cell(row=3, column=1, value=("【下書き】" if draft else "") + _NOTICE).font = Font(color="A8352A")

    header = 5
    for column, label in enumerate(_COLUMNS, start=1):
        cell = sheet.cell(row=header, column=column, value=label)
        cell.font = Font(bold=True)
    for offset, monday in enumerate(weeks):
        cell = sheet.cell(row=header, column=len(_COLUMNS) + 1 + offset, value=monday.strftime("%m/%d"))
        cell.font = Font(bold=True, size=8)
        cell.alignment = Alignment(textRotation=90)
        sheet.column_dimensions[get_column_letter(len(_COLUMNS) + 1 + offset)].width = 3.2

    fills = {name: PatternFill("solid", fgColor=color) for name, color in _FILL.items()}
    for index, row in enumerate(wbs.walk(), start=header + 1):
        _write_row(sheet, index, row, weeks, fills)

    for column, width in enumerate((7, 40, 12, 14, 8, 11, 11, 6, 11, 11, 7), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = sheet.cell(row=header + 1, column=len(_COLUMNS) + 1)
    outline = sheet.sheet_properties.outlinePr
    if outline is not None:
        outline.summaryBelow = False  # 親を上に置く（WBS の並びと合わせる）
    book.save(path)
