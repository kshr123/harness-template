"""出力形式の登録簿と、表計算の写しのテスト。

既定が 1 形式のままであること（成果物が勝手に 2 つ並ばない）と、増やした形式が正本と食い違わないこと
（書いた表を読み戻して木から導いた値と突き合わせる）を見る。
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

import frontmatter
import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from harness.deliver import formats
from harness.deliver import wbs as wbs_mod
from harness.deliver.cli import wbs_app
from harness.deliver.overlay import Overlay

pytestmark = pytest.mark.integration

TODAY = date(2026, 8, 20)


def _write(path: Path, meta: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    post = frontmatter.Post("")
    post.metadata.update(meta)
    path.write_text(frontmatter.dumps(post), encoding="utf-8")


def _scaffold(root: Path) -> None:
    alpha = root / "work" / "EP-90-alpha"
    _write(alpha / "item.md", {"id": "EP-90", "kind": "epic", "status": "in-progress", "plan": "detailed"})
    _write(
        alpha / "T-9001-a.md",
        {"id": "T-9001", "kind": "task", "status": "done", "title": "設計", "start": "2026-08-03", "due": "2026-08-07"},
    )
    _write(
        alpha / "T-9002-b.md",
        {"id": "T-9002", "kind": "task", "status": "todo", "title": "実装", "start": "2026-08-10", "due": "2026-08-12"},
    )


def test_html_is_always_available(tmp_path: Path) -> None:
    assert "html" in formats.RENDERERS


def test_the_spreadsheet_branch_appears_only_with_its_dependency() -> None:
    """依存が入っている環境では一覧に出る（入れていなければ出ない＝選べる形式が使える形式）。"""
    assert "xlsx" in formats.RENDERERS  # 開発環境は全部入り（uv sync --all-extras）
    assert formats.RENDERERS.extras_hint["xlsx"] == "openpyxl"


def test_every_registered_format_explains_itself() -> None:
    """一覧に出る以上、説明文がある（登録簿が説明文を必須にしている）。"""
    for kind in formats.RENDERERS:
        assert formats.RENDERERS[kind].description.strip()


def test_an_unknown_format_fails_with_the_candidates(tmp_path: Path) -> None:
    _scaffold(tmp_path)
    result = CliRunner().invoke(
        wbs_app, ["export", "--root", str(tmp_path), "--format", "pdf", "--today", TODAY.isoformat()]
    )
    assert result.exit_code == 1
    assert "html" in result.output  # 候補を出す


def test_the_default_format_is_html_and_the_suffix_follows_it(tmp_path: Path) -> None:
    """形式を指定しなければ HTML 1 つだけが出る（成果物が勝手に 2 つ並ばない）。"""
    _scaffold(tmp_path)
    result = CliRunner().invoke(wbs_app, ["export", "--root", str(tmp_path), "--today", TODAY.isoformat()])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "artifacts" / "wbs" / "WBS.html").is_file()
    assert not (tmp_path / "artifacts" / "wbs" / "WBS.xlsx").exists()


def test_the_spreadsheet_matches_the_tree(tmp_path: Path) -> None:
    """書いた表を読み戻し、木から導いた値と突き合わせる（書式の崩れを人の目に頼らない）。"""
    _scaffold(tmp_path)
    out = tmp_path / "WBS.xlsx"
    result = CliRunner().invoke(
        wbs_app,
        ["export", "--root", str(tmp_path), "--out", str(out), "--format", "xlsx", "--today", TODAY.isoformat()],
    )
    assert result.exit_code == 0, result.output

    built = wbs_mod.build(tmp_path, today=TODAY, overlay=Overlay())
    sheet = load_workbook(out).active
    assert sheet is not None
    header = 5

    def _as_date(value: object) -> date | None:
        """表計算のセルから日付を取り出す（日付として書けていなければ None＝突き合わせで落ちる）。"""
        return value.date() if isinstance(value, datetime) else None

    read = {
        str(sheet.cell(row=header + i + 1, column=1).value): (
            _as_date(sheet.cell(row=header + i + 1, column=6).value),
            _as_date(sheet.cell(row=header + i + 1, column=7).value),
            sheet.cell(row=header + i + 1, column=8).value,
        )
        for i in range(len(built.walk()))
    }
    for row in built.walk():
        start, due, days = read[row.code]
        assert start == row.start
        assert due == row.due
        assert days == row.workdays


def test_the_spreadsheet_says_it_is_a_copy(tmp_path: Path) -> None:
    """編集を誘う道具なので、取り込まれない写しであることを本文に書く（黙って捨てない）。"""
    _scaffold(tmp_path)
    out = tmp_path / "WBS.xlsx"
    CliRunner().invoke(
        wbs_app,
        ["export", "--root", str(tmp_path), "--out", str(out), "--format", "xlsx", "--today", TODAY.isoformat()],
    )
    sheet = load_workbook(out).active
    assert sheet is not None
    assert "取り込まれません" in str(sheet.cell(row=3, column=1).value)


def test_against_is_refused_for_non_html_and_with_at(tmp_path: Path) -> None:
    """--against（ベースライン重ね描き）は HTML のみ・--at とは排他（黙って層を落とさず明示的に拒否）。"""
    _scaffold(tmp_path)
    runner = CliRunner()
    out = tmp_path / "WBS.xlsx"
    r_xlsx = runner.invoke(
        wbs_app,
        [
            "export",
            "--root",
            str(tmp_path),
            "--against",
            "HEAD",
            "--format",
            "xlsx",
            "--out",
            str(out),
            "--today",
            TODAY.isoformat(),
        ],
    )
    assert r_xlsx.exit_code == 1 and "HTML のみ" in r_xlsx.output and not out.exists()
    r_at = runner.invoke(
        wbs_app,
        [
            "export",
            "--root",
            str(tmp_path),
            "--against",
            "HEAD",
            "--at",
            "HEAD",
            "--out",
            str(tmp_path / "WBS.html"),
            "--today",
            TODAY.isoformat(),
        ],
    )
    assert r_at.exit_code == 1 and "同時に使えない" in r_at.output


def test_the_spreadsheet_does_not_write_live_formulas_from_user_text(tmp_path: Path) -> None:
    """作業名・チーム・担当が数式記号（=+-@）で始まっても、生きた数式にせず文字として書く（数式注入対策）。

    写しはクライアントに渡す読み取り専用ファイル。openpyxl は先頭 `=` の文字列を数式セル（data_type 'f'）
    として保存してしまうので、そのままだと先方の Excel で任意の式が走る。文字（'s'）で書けていることを見る。
    """
    alpha = tmp_path / "work" / "EP-90-alpha"
    _write(alpha / "item.md", {"id": "EP-90", "kind": "epic", "status": "todo", "plan": "detailed"})
    _write(
        alpha / "T-9001-a.md",
        {
            "id": "T-9001",
            "kind": "task",
            "status": "todo",
            "title": '=HYPERLINK("http://evil.example","x")',
            "team": "=1+1",
            "start": "2026-08-03",
            "due": "2026-08-07",
        },
    )
    out = tmp_path / "WBS.xlsx"
    result = CliRunner().invoke(
        wbs_app,
        ["export", "--root", str(tmp_path), "--out", str(out), "--format", "xlsx", "--today", TODAY.isoformat()],
    )
    assert result.exit_code == 0, result.output
    sheet = load_workbook(out).active
    assert sheet is not None
    # 作業名（2 列目）・チーム（3 列目）が数式（'f'）でなく文字（'s'）で書かれている。
    name_cell = next(c for r in sheet.iter_rows() for c in r if isinstance(c.value, str) and "HYPERLINK" in c.value)
    assert name_cell.data_type == "s", "作業名が生きた数式として書かれている（数式注入）"
    team_cell = next(c for r in sheet.iter_rows() for c in r if c.value == "=1+1")
    assert team_cell.data_type == "s", "チームが生きた数式として書かれている（数式注入）"


def test_the_spreadsheet_marks_the_today_week_and_holiday_weeks_and_has_a_legend(tmp_path: Path) -> None:
    """xlsx の体裁パリティ：本日を含む週は見出しを赤字・休業（祝日/会社休）を含む週は淡い灰・凡例行を出す。"""
    from openpyxl import load_workbook

    from harness.deliver import wbs as wbs_mod
    from harness.deliver.overlay import Overlay
    from harness.deliver.render import COLUMN_LABELS
    from harness.deliver.xlsx import write_xlsx

    _write(
        tmp_path / "work" / "T-1.md",
        {"id": "T-0001", "kind": "task", "status": "todo", "start": "2026-08-03", "due": "2026-08-28"},
    )
    # 会社休を 08-19（水）に置く＝08/17 の週が休業週。本日は別の週（08-06）に置いて印を分けて確かめる。
    overlay = Overlay.model_validate({"calendar": {"extra_holidays": ["2026-08-19"]}})
    built = wbs_mod.build(tmp_path, today=date(2026, 8, 6), overlay=overlay)
    out = tmp_path / "WBS.xlsx"
    write_xlsx(built, out)
    sheet = load_workbook(out).active
    assert sheet is not None
    base = len(COLUMN_LABELS) + 1
    weeks = {sheet.cell(row=5, column=c).value: sheet.cell(row=5, column=c) for c in range(base, base + 8)}
    # 本日（08-06）は 08/03 の週＝赤字。
    assert weeks["08/03"].font.color.rgb.endswith("B12F1F")
    # 休業を含む 08/17 の週＝淡い灰の塗り。休業の無い 08/03 の週は塗らない。
    assert weeks["08/17"].fill.patternType == "solid" and weeks["08/17"].fill.fgColor.rgb.endswith("EDF0F3")
    assert weeks["08/03"].fill.patternType is None
    # 凡例行（意味の対応）が出る。
    assert sheet.cell(row=4, column=1).value == "凡例"
    assert {sheet.cell(row=4, column=c).value for c in range(2, 8)} >= {"期間", "完了", "遅れ", "休業週", "本日"}


def test_the_spreadsheet_folds_by_depth(tmp_path: Path) -> None:
    """階層は表計算側の折りたたみ（アウトライン）で表す＝先方が畳んで読める。"""
    _scaffold(tmp_path)
    out = tmp_path / "WBS.xlsx"
    CliRunner().invoke(
        wbs_app,
        ["export", "--root", str(tmp_path), "--out", str(out), "--format", "xlsx", "--today", TODAY.isoformat()],
    )
    sheet = load_workbook(out).active
    assert sheet is not None
    assert sheet.row_dimensions[6].outlineLevel == 0  # 親（1）
    assert sheet.row_dimensions[7].outlineLevel == 1  # 子（1.1）
